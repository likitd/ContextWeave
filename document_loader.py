import os
from pathlib import Path
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
import config


def load_file_as_text(file_path: str) -> str | None:
    ext = Path(file_path).suffix.lower()

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if ext == ".docx":
        from docx import Document as DocxDocument
        doc = DocxDocument(file_path)
        return "\n".join(para.text for para in doc.paragraphs)

    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(file_path)
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    texts.append(shape.text)
        return "\n".join(texts)

    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(file_path, read_only=True, data_only=True)
        rows = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            rows.append(f"--- Sheet: {sheet} ---")
            for row in ws.iter_rows(values_only=True):
                rows.append("\t".join(str(cell) if cell is not None else "" for cell in row))
        return "\n".join(rows)

    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception:
        return None


def discover_files(folder: str) -> list[str]:
    files = []
    for root, _, filenames in os.walk(folder):
        for fname in filenames:
            ext = Path(fname).suffix.lower()
            if ext in config.SUPPORTED_EXTENSIONS:
                files.append(os.path.join(root, fname))
    return sorted(files)


def load_documents(folder: str = None) -> list[Document]:
    folder = folder or config.CONTEXT_FOLDER
    file_paths = discover_files(folder)
    documents = []

    for fp in file_paths:
        content = load_file_as_text(fp)
        if not content or not content.strip():
            continue
        rel_path = os.path.relpath(fp, folder)
        documents.append(Document(
            page_content=content,
            metadata={
                "source": rel_path,
                "absolute_path": fp,
                "extension": Path(fp).suffix.lower(),
            }
        ))

    print(f"Loaded {len(documents)} documents from {folder}")
    return documents


def split_documents(documents: list[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=config.CHUNK_SIZE,
        chunk_overlap=config.CHUNK_OVERLAP,
        length_function=len,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(documents)
    print(f"Split into {len(chunks)} chunks")
    return chunks
